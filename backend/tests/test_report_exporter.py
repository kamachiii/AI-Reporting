"""Unit test untuk layanan ekspor Excel berformat + native chart (report_exporter.py)."""
import io
import openpyxl
import pytest
from app.services.report_exporter import generate_excel_report, _parse_numeric, _is_currency_column


def test_parse_numeric_indonesian_and_raw():
    assert _parse_numeric(100) == (True, 100.0)
    assert _parse_numeric(25.5) == (True, 25.5)
    assert _parse_numeric("Rp 14.563.500") == (True, 14563500.0)
    assert _parse_numeric("14.563.500,50") == (True, 14563500.5)
    assert _parse_numeric("73") == (True, 73.0)
    assert _parse_numeric("FOUR BEST SYNERGY")[0] is False
    assert _parse_numeric(None)[0] is False


def test_is_currency_column():
    assert _is_currency_column("total_omzet") is True
    assert _is_currency_column("harga_jual") is True
    assert _is_currency_column("ar_leasing") is True
    assert _is_currency_column("nilai_transaksi") is True
    assert _is_currency_column("total_nominal") is True
    assert _is_currency_column("jumlah_unit") is False
    assert _is_currency_column("kuartal") is False
    assert _is_currency_column("kuantiti_part_terjual") is False
    assert _is_currency_column("total_transaksi") is False
    assert _is_currency_column("total_pkb") is False


def _get_chart_title_text(chart) -> str:
    """Helper untuk mengambil teks title dari openpyxl chart title object."""
    if hasattr(chart.title, 'tx') and chart.title.tx and chart.title.tx.rich:
        return chart.title.tx.rich.p[0].r[0].t
    return str(chart.title)


def test_generate_excel_with_time_series_chart():
    rows = [
        {"kuartal": "Q1", "unit_terjual": 129, "total_omzet": 25735500000},
        {"kuartal": "Q2", "unit_terjual": 73, "total_omzet": 14563500000},
        {"kuartal": "Q3", "unit_terjual": 83, "total_omzet": 16558500000},
        {"kuartal": "Q4", "unit_terjual": 65, "total_omzet": 12967500000},
    ]
    columns = ["kuartal", "unit_terjual", "total_omzet"]
    excel_bytes = generate_excel_report(
        question="performa penjualan per kuartal 2025",
        branch_code="TST_01",
        tab_name="Unit Kendaraan",
        rows=rows,
        columns=columns
    )
    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 2000

    # Verifikasi dapat dibaca openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    ws = wb.active
    assert ws.title == "Laporan Eksekutif"
    assert ws['A1'].value == "DMS AI PLATFORM — LAPORAN EKSEKUTIF"
    # Memeriksa baris data terisi
    assert ws['A5'].value == "Kuartal"
    assert ws['B5'].value == "Unit Terjual"
    assert ws['C5'].value == "Total Omzet"
    assert ws['A6'].value == "Q1"
    assert ws['B6'].value == 129.0
    assert ws['C6'].value == 25735500000.0

    # Memeriksa chart tersemat
    assert len(ws._charts) > 0
    chart = ws._charts[0]
    assert "Unit Kendaraan" in _get_chart_title_text(chart)


def test_generate_excel_categorical_comparison():
    rows = [
        {"customer": "PT Surya", "total_pembelian": 1500000000},
        {"customer": "CV Maju", "total_pembelian": 850000000},
        {"customer": "Bpk Budi", "total_pembelian": 450000000},
    ]
    columns = ["customer", "total_pembelian"]
    excel_bytes = generate_excel_report(
        question="top 3 customer 2025",
        branch_code="TST_01",
        tab_name=None,
        rows=rows,
        columns=columns
    )
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    ws = wb.active
    assert len(ws._charts) > 0
    assert _get_chart_title_text(ws._charts[0]) == "Grafik Perbandingan"


def test_generate_excel_empty_rows():
    excel_bytes = generate_excel_report(
        question="cek data kosong",
        branch_code="TST_01",
        tab_name=None,
        rows=[],
        columns=["id", "nama"]
    )
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    ws = wb.active
    assert ws['A5'].value == "Id"
    assert ws['B5'].value == "Nama"


def test_endpoint_export_excel():
    from fastapi.testclient import TestClient
    from app.main import app
    from app.core.security import require_user_role

    client = TestClient(app)
    app.dependency_overrides[require_user_role] = lambda: {
        "user_id": 99,
        "username": "tester01",
        "role": "user",
        "allowed_branches": ["TST_01"]
    }

    payload = {
        "branch_code": "TST_01",
        "question": "penjualan kuartal 2025",
        "tab_name": "Unit Kendaraan",
        "rows": [{"kuartal": "Q1", "omzet": 50000000}],
        "columns": ["kuartal", "omzet"]
    }
    response = client.post("/chat/export-excel", json=payload)
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "attachment; filename=" in response.headers["content-disposition"]
    assert len(response.content) > 1000

    # Cabang tidak diizinkan -> 403
    payload["branch_code"] = "CABANG_LAIN"
    bad_resp = client.post("/chat/export-excel", json=payload)
    assert bad_resp.status_code == 403

    # Uji coba payload dengan rows berupa list of lists
    payload_lists = {
        "branch_code": "TST_01",
        "question": "penjualan kuartal 2025",
        "tab_name": "Unit Kendaraan",
        "rows": [[1, 129, 25735500000, 199500000], [2, 73, 14563500000, 199500000]],
        "columns": ["kuartal", "jumlah_unit_terjual", "total_omzet_penjualan", "rata_rata_harga_jual_unit"]
    }
    resp_lists = client.post("/chat/export-excel", json=payload_lists)
    assert resp_lists.status_code == 200
    assert len(resp_lists.content) > 1000

    app.dependency_overrides.clear()
